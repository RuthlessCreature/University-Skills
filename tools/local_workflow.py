"""Bounded file workflow for teaching. No model calls or claims of autonomous research.

Supports structured JSON intake + native DOCX template. Does not parse arbitrary PDF/OCR.
"""
import json,pathlib,hashlib,datetime,copy
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

class WorkflowError(ValueError):pass
def digest(p):return hashlib.sha256(pathlib.Path(p).read_bytes()).hexdigest()
def dump(p,v):pathlib.Path(p).write_text(json.dumps(v,ensure_ascii=False,indent=2),'utf-8')
def load(p):return json.loads(pathlib.Path(p).read_text('utf-8-sig'))
def validate(m):
    if m.get('kind') not in ['grant','paper','patent']:raise WorkflowError('unsupported_kind')
    if m.get('material_type') not in ['teaching-simulation','public-source','authorized-user']:raise WorkflowError('material_type_required')
    if not m.get('sections') or not isinstance(m['sections'],list):raise WorkflowError('sections_required')
    titles=[x['title'] for x in m['sections']]
    if len(set(titles))!=len(titles):raise WorkflowError('duplicate_section')
    sources={s['id']:s for s in m.get('sources',[])}
    if len(sources)!=len(m.get('sources',[])):raise WorkflowError('duplicate_source')
    for s in sources.values():
        if not s.get('locator') or not s.get('type'):raise WorkflowError('source_locator_required')
    for sec in m['sections']:
        if not isinstance(sec.get('content'),str):raise WorkflowError('content_required')
        if len(sec['content'])>sec.get('max_chars',5000):raise WorkflowError('section_length_exceeded:'+sec['title'])
        for sid in sec.get('source_ids',[]):
            if sid not in sources:raise WorkflowError('unknown_source:'+sid)
    goals={g['id'] for g in m.get('goals',[])}; reqs={r['id'] for r in m.get('requirements',[])}
    for r in m.get('requirements',[]):
        if not r.get('goal_ids') or not set(r['goal_ids'])<=goals:raise WorkflowError('unmapped_requirement:'+r['id'])
        if not r.get('artifact') or not r.get('verification'):raise WorkflowError('unverifiable_requirement:'+r['id'])
    if goals and any(not any(g in r['goal_ids'] for r in m.get('requirements',[])) for g in goals):raise WorkflowError('unmapped_goal')
    return sources
def make_template(p,sections):
    d=Document();style=d.styles['Normal'];style.font.name='Microsoft YaHei';style.font.size=Pt(11)
    style._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');style.paragraph_format.space_after=Pt(8)
    for sn in ['Title','Heading 1','Heading 2']:
        d.styles[sn].font.name='Microsoft YaHei';d.styles[sn].font.color.rgb=RGBColor(0,0,0);d.styles[sn]._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    sec=d.sections[0];sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=sec.right_margin=Cm(2.2)
    d.add_paragraph('科研教学文档模板',style='Title')
    for title in sections:d.add_paragraph(title,style='Heading 1');d.add_paragraph('{{'+title+'}}')
    for tree in [d._element,d.styles.element]:
        for el in list(tree.iter(qn('w:pBdr'))):el.getparent().remove(el)
    d.save(p)
def run(input_path,template_path,out,approval_path=None,final=False,guide_path=None):
    out=pathlib.Path(out);out.mkdir(parents=True,exist_ok=True); events=[]
    def event(stage,state,detail=''):events.append(dict(stage=stage,state=state,detail=detail,utc=datetime.datetime.now(datetime.timezone.utc).isoformat()))
    event('intake','running');m=load(input_path);input_hash=digest(input_path);template_hash=digest(template_path)
    try:
        guide_hash=None
        if m.get('kind')=='grant':
            if not guide_path:raise WorkflowError('guide_required')
            guide=load(guide_path);guide_hash=digest(guide_path)
            if guide.get('type')!=m.get('material_type'):raise WorkflowError('guide_material_type_mismatch')
            if [s['title'] for s in m['sections']]!=guide.get('required_sections'):raise WorkflowError('guide_section_mismatch')
            if m.get('requirements')!=guide.get('requirements'):raise WorkflowError('guide_requirement_mismatch')
        sources=validate(m);event('constraints','passed')
        d=Document(template_path); present=[p.text for p in d.paragraphs]
        for sec in m['sections']:
            if sec['title'] not in present or '{{'+sec['title']+'}}' not in present:raise WorkflowError('template_missing_section:'+sec['title'])
        approval=load(approval_path) if approval_path else None
        approved=bool(approval and approval.get('decision')=='approved' and approval.get('reviewer_type')=='human' and approval.get('reviewer') and approval.get('input_sha256')==input_hash and approval.get('template_sha256')==template_hash and approval.get('guide_sha256')==guide_hash)
        if approval and not approved:raise WorkflowError('invalid_or_stale_human_approval')
        issues=copy.deepcopy(m.get('issues',[]))
        for sec in m['sections']:
            if sec.get('factual',True) and not sec.get('source_ids'):issues.append(dict(id='SOURCE-'+sec['title'],severity='blocking',issue='该节没有来源定位'))
        if not approved:issues.append(dict(id='HUMAN-SIGNOFF',severity='blocking',issue='需要研究者对输入和模板的当前版本签核'))
        if final and (not approved or any(i.get('severity')=='blocking' for i in issues)):raise WorkflowError('final_export_blocked')
        event('outline','approved' if approved else 'human_review_required')
        # Fill existing template paragraphs; no invented research prose is generated.
        for p in d.paragraphs:
            if p.style.name=='Title':p.text=m['title']+' 教学草稿'
            for sec in m['sections']:
                if p.text=='{{'+sec['title']+'}}':p.text=sec['content']+'\n依据：'+', '.join(sec.get('source_ids',[]))
        d.add_paragraph('来源与未决问题',style='Heading 1')
        d.add_paragraph('材料属性：'+m['material_type']+'。本文件是受约束的材料组织草稿。科研事实与正式提交仍需研究者核对。')
        for sid,s in sources.items():d.add_paragraph(sid+'  '+s['locator'])
        for i in issues:d.add_paragraph(i['id']+'  '+i['issue'])
        for tree in [d._element,d.styles.element]:
            for el in list(tree.iter(qn('w:pBdr'))):el.getparent().remove(el)
        artifact=out/(m['kind']+'-draft.docx');d.save(artifact);event('docx','written',artifact.name)
        mapping=m.get('requirements',[]);dump(out/'requirements.json',mapping);dump(out/'mapping.json',dict(goals=m.get('goals',[]),requirements=mapping))
        dump(out/'evidence-ledger.json',{'scope':'source identity and supplied locator only; scientific truth not auto-verified','sources':list(sources.values()),'sections':[{'title':s['title'],'source_ids':s.get('source_ids',[])} for s in m['sections']]})
        dump(out/'issues.json',issues);dump(out/'questions.json',m.get('questions',[]))
        route='flowchart TD\n';steps=m.get('route',['输入材料','核对事实','草稿组织','研究者复核'])
        for i,x in enumerate(steps):route+=f'  S{i}["{x.replace(chr(34),chr(39))}"]\n'
        for i in range(len(steps)-1):route+=f'  S{i} --> S{i+1}\n'
        (out/'route.mmd').write_text(route,'utf-8')
        result=dict(kind=m['kind'],status='final_ready' if final else 'draft_review_required',input_sha256=input_hash,template_sha256=template_hash,guide_sha256=guide_hash,artifact=artifact.name,artifact_sha256=digest(artifact),human_approved=approved,external_model_called=False,scientific_validation='not_performed',events=events)
        dump(out/'run.json',result);return result
    except Exception as e:
        event('workflow','failed',str(e));dump(out/'run.json',dict(status='failed',error=str(e),input_sha256=input_hash,template_sha256=template_hash,events=events));raise
