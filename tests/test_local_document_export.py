"""Real Word/package and artifact handoff tests, with no provider simulation."""
from pathlib import Path
from zipfile import ZipFile
import json,sys,tempfile,unittest
from docx import Document
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'tools'))
from local_skill_runner import LocalTools,sha,validate_resume
from local_document_export import inspect_reference

class WordExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.reference=self.root/'reference.docx'
        d=Document();d.add_paragraph('Protected heading');d.add_paragraph('');d.add_paragraph('Protected instructions')
        d.sections[0].footer.paragraphs[0].text='Fixture footer'
        d.add_table(rows=1,cols=2).cell(0,0).text='Native fixture table';d.save(self.reference)
        inventory=inspect_reference(self.reference)
        self.profile=self.root/'profile.json'
        self.definition={'mode':'ooxml-slots','reference_alias':'template','reference_sha256':sha(self.reference),'slots':[{'id':'body','paragraph_index':1,'anchor_xml_sha256':inventory['paragraphs'][1]['xml_sha256'],'required':True}],'protected_paragraph_indices':[0,2]}
        self.profile.write_text(json.dumps(self.definition),encoding='utf-8')
        self.local=LocalTools({'template':self.reference,'profile':self.profile},self.root/'out')
        self.payload={'material_type':'teaching-simulation','slots':{'body':['课程样例。拟开展工作尚未执行。']},'source_notes':['Fixture source'],'unresolved_questions':['申请人信息待提供。']}
    def tearDown(self):self.tmp.cleanup()
    def export(self,payload=None,name='draft.docx'):
        self.local.save('content.json',json.dumps(payload or self.payload,ensure_ascii=False))
        return self.local.invoke('export_word',{'name':name,'profile_alias':'profile','content_name':'content.json'})
    def test_actual_editable_word_and_source_preservation(self):
        original=sha(self.reference);result=self.export();p=Path(result['path']);d=Document(p)
        self.assertIn('课程样例。拟开展工作尚未执行。',[x.text for x in d.paragraphs]);self.assertEqual(sha(self.reference),original)
        self.assertEqual(d.tables[0].cell(0,0).text,'Native fixture table')
        self.assertEqual(d.sections[0].footer.paragraphs[0].text,'Fixture footer')
        self.assertIn('申请人信息待提供。',''.join(x.text for x in d.paragraphs))
        with ZipFile(self.reference) as a,ZipFile(p) as b:
            for n in a.namelist():
                if n!='word/document.xml':self.assertEqual(a.read(n),b.read(n),n)
        self.assertEqual(result['scientific_review'],'NOT_PERFORMED')
    def test_unknown_slot_rejected_without_word(self):
        self.payload['slots']['invented']=['Unrequested text']
        with self.assertRaises(ValueError):self.export()
        self.assertFalse((self.local.output/'draft.docx').exists())
    def test_missing_slot_rejected(self):
        self.payload['slots']={}
        with self.assertRaises(ValueError):self.export()
    def test_unknown_payload_not_silently_dropped(self):
        self.payload['extra_research_result']='An unsupported additional result'
        with self.assertRaises(ValueError):self.export()
    def test_wrong_template_anchor_rejected(self):
        self.definition['slots'][0]['anchor_xml_sha256']='0'*64;self.profile.write_text(json.dumps(self.definition),encoding='utf-8')
        self.local=LocalTools({'template':self.reference,'profile':self.profile},self.root/'out')
        with self.assertRaises(ValueError):self.export()
    def test_protected_instruction_removal_rejected(self):
        self.definition['remove_guidance_indices']=[2];self.profile.write_text(json.dumps(self.definition),encoding='utf-8')
        self.local=LocalTools({'template':self.reference,'profile':self.profile},self.root/'out')
        with self.assertRaises(ValueError):self.export()
    def test_no_word_overwrite(self):
        result=self.export();before=sha(Path(result['path']))
        with self.assertRaises(FileExistsError):self.local.invoke('export_word',{'name':'draft.docx','profile_alias':'profile','content_name':'content.json'})
        self.assertEqual(sha(Path(result['path'])),before)
    def test_course_document_has_native_table_and_clear_status(self):
        p=self.root/'course-profile.json';p.write_text(json.dumps({'mode':'course-document'}),encoding='utf-8')
        local=LocalTools({'profile':p},self.root/'course-out')
        local.save('paper.json',json.dumps({'title':'公开数据课程示例','sections':[{'heading':'结果','paragraphs':['以下为测试输入。']}],'tables':[{'columns':['组','n'],'rows':[['A',2]]}],'unresolved_questions':['作者待确认。']},ensure_ascii=False))
        result=local.invoke('export_word',{'name':'paper.docx','profile_alias':'profile','content_name':'paper.json'})
        d=Document(result['path']);self.assertEqual(len(d.tables),1);self.assertEqual(d.tables[0].cell(1,1).text,'2')
        self.assertIn('COURSE_WORKING_DOCUMENT',result['detail']['template_status'])
    def test_render_refuses_unexported_or_changed_file(self):
        with self.assertRaises(ValueError):self.local.invoke('render_word',{'name':'arbitrary.docx'})
        self.export();(self.local.output/'draft.docx').write_bytes(b'changed')
        with self.assertRaises(ValueError):self.local.invoke('render_word',{'name':'draft.docx'})

class SectionHandoffTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.profile=self.root/'profile.json'
        self.profile.write_text(json.dumps({'mode':'ooxml-slots','slots':[{'id':'body','required':True}]}),encoding='utf-8')
        self.part=self.root/'part.json';self.part.write_text(json.dumps({'slots':{'body':['实际文件中的段落。']}}),encoding='utf-8')
        self.local=LocalTools({'profile':self.profile,'part':self.part},self.root/'out')
    def tearDown(self):self.tmp.cleanup()
    def test_assembly_reads_actual_artifact(self):
        r=self.local.invoke('assemble_word_content',{'name':'merged.json','profile_alias':'profile','section_aliases':['part']})
        self.assertEqual(json.loads(Path(r['path']).read_text('utf-8'))['slots']['body'],['实际文件中的段落。'])
    def test_duplicate_slot_requires_explicit_version_choice(self):
        with self.assertRaises(ValueError):self.local.invoke('assemble_word_content',{'name':'merged.json','profile_alias':'profile','section_aliases':['part','part']})
        self.assertFalse((self.local.output/'merged.json').exists())
    def test_previous_node_becomes_readable_and_remains_hash_guarded(self):
        checkpoint=self.root/'node.json';checkpoint.write_text(json.dumps({'artifacts':[{'name':'part.json','sha256':sha(self.part)}]}),encoding='utf-8')
        nodes=[{'node':'background','checkpoint':str(checkpoint),'sha256':sha(checkpoint)}]
        aliases=self.local.register_previous(nodes,self.root)
        self.assertEqual(json.loads(self.local.read(aliases[0])['text'])['slots']['body'],['实际文件中的段落。'])
        self.part.write_text('modified',encoding='utf-8')
        with self.assertRaises(ValueError):self.local.read(aliases[0])

if __name__=='__main__':unittest.main(verbosity=2)
